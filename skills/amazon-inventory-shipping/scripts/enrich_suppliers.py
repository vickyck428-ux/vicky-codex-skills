#!/usr/bin/env python3

import argparse
import csv
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook


SUPPLIER_COLUMNS = [
    "supplier_name",
    "supplier_1688_url",
    "supplier_price",
    "supplier_notes",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge supplier matches into an existing shipment plan."
    )
    parser.add_argument("--plan", required=True, help="Shipment plan CSV or XLSX file.")
    parser.add_argument(
        "--suppliers",
        required=True,
        help="Supplier match CSV or XLSX file keyed by sku or product_name.",
    )
    parser.add_argument("--output", required=True, help="Merged output CSV or XLSX file.")
    return parser.parse_args()


def normalize_header(value: str) -> str:
    return value.strip()


def read_csv(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError(f"{path} has no header row.")
        headers = [normalize_header(name) for name in reader.fieldnames]
        rows = []
        for raw_row in reader:
            row = {}
            for header, value in zip(headers, raw_row.values()):
                row[header] = value
            rows.append(row)
    return headers, rows


def read_xlsx(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"{path} is empty.")
    headers = [normalize_header(str(cell or "")) for cell in values[0]]
    rows = []
    for raw_row in values[1:]:
        row = {}
        for index, header in enumerate(headers):
            row[header] = raw_row[index] if index < len(raw_row) else None
        rows.append(row)
    return headers, rows


def read_rows(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError("Only CSV and XLSX are supported.")


def write_csv(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def write_xlsx(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)


def write_rows(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        write_csv(path, headers, rows)
        return
    if suffix in {".xlsx", ".xlsm"}:
        write_xlsx(path, headers, rows)
        return
    raise ValueError("Only CSV and XLSX are supported.")


def value_key(row: Dict[str, object], field: str) -> str:
    value = row.get(field)
    if value is None:
        return ""
    return str(value).strip().lower()


def build_supplier_index(rows: List[Dict[str, object]]) -> Dict[str, Dict[str, object]]:
    index: Dict[str, Dict[str, object]] = {}
    for row in rows:
        for field in ("sku", "product_name"):
            key = value_key(row, field)
            if key and key not in index:
                index[key] = row
    return index


def main() -> None:
    args = parse_args()
    plan_path = Path(args.plan).expanduser().resolve()
    suppliers_path = Path(args.suppliers).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()

    plan_headers, plan_rows = read_rows(plan_path)
    _, supplier_rows = read_rows(suppliers_path)
    supplier_index = build_supplier_index(supplier_rows)

    output_rows: List[Dict[str, object]] = []
    for row in plan_rows:
        supplier_row = None
        for field in ("sku", "product_name"):
            key = value_key(row, field)
            if key and key in supplier_index:
                supplier_row = supplier_index[key]
                break

        merged = dict(row)
        for column in SUPPLIER_COLUMNS:
            merged[column] = supplier_row.get(column, "") if supplier_row else ""
        output_rows.append(merged)

    output_headers = list(plan_headers)
    for column in SUPPLIER_COLUMNS:
        if column not in output_headers:
            output_headers.append(column)

    write_rows(output_path, output_headers, output_rows)
    print(f"Wrote {len(output_rows)} rows to {output_path}")


if __name__ == "__main__":
    main()
