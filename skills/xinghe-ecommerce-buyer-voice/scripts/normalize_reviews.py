#!/usr/bin/env python3
"""Normalize ecommerce review exports without inventing semantic labels.

Supports CSV, TSV, JSON, JSONL, TXT and XLSX (when openpyxl is available).
Writes normalized_reviews.csv and normalization_summary.json.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


FIELDS = [
    "platform", "source_id", "source_url", "product", "sku", "date", "rating",
    "review_text", "buyer_signal", "scenario", "purchase_motivation",
    "purchase_concern", "satisfaction", "dissatisfaction", "return_reason",
    "issue_owner", "severity", "commercial_impact", "duplicate_status",
]

ALIASES = {
    "platform": ["platform", "平台", "站点", "渠道"],
    "source_id": ["source_id", "review_id", "order_id", "id", "评价id", "评论id", "订单号"],
    "source_url": ["source_url", "url", "link", "链接", "来源链接"],
    "product": ["product", "product_name", "商品", "商品名称", "产品", "产品名称"],
    "sku": ["sku", "variant", "variation", "规格", "型号", "款式", "颜色尺码"],
    "date": ["date", "review_date", "created_at", "日期", "评价时间", "评论时间", "时间"],
    "rating": ["rating", "stars", "score", "星级", "评分", "打分"],
    "review_text": ["review_text", "review", "content", "comment", "text", "评价内容", "评论内容", "内容", "买家原话"],
    "buyer_signal": ["buyer_signal", "buyer", "买家", "用户", "人群信号"],
    "scenario": ["scenario", "use_case", "使用场景", "场景"],
    "purchase_motivation": ["purchase_motivation", "motivation", "购买动机", "购买理由"],
    "purchase_concern": ["purchase_concern", "concern", "购买顾虑", "顾虑"],
    "satisfaction": ["satisfaction", "like", "满意点", "喜欢"],
    "dissatisfaction": ["dissatisfaction", "dislike", "不满意点", "缺点"],
    "return_reason": ["return_reason", "退款原因", "退货原因"],
    "issue_owner": ["issue_owner", "问题归属", "问题类型"],
    "severity": ["severity", "严重度", "情绪强度"],
    "commercial_impact": ["commercial_impact", "商业影响", "转化影响"],
}


def norm_header(value: Any) -> str:
    return re.sub(r"[\s_\-—/\\:：()（）]+", "", str(value or "").strip().lower())


def norm_text(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def dedupe_text(value: str) -> str:
    value = value.lower()
    value = re.sub(r"\s+", "", value)
    value = re.sub(r"[^\w\u4e00-\u9fff]", "", value)
    return value


def decode_csv(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码；请另存为 UTF-8 或 GB18030。")


def read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        text = decode_csv(path)
        delimiter = "\t" if suffix == ".tsv" else ","
        return [dict(row) for row in csv.DictReader(text.splitlines(), delimiter=delimiter)]
    if suffix == ".json":
        data = json.loads(decode_csv(path))
        if isinstance(data, dict):
            for key in ("reviews", "data", "items", "records"):
                if isinstance(data.get(key), list):
                    data = data[key]
                    break
        if not isinstance(data, list):
            raise ValueError("JSON 顶层应为数组，或包含 reviews/data/items/records 数组。")
        return [item if isinstance(item, dict) else {"review_text": item} for item in data]
    if suffix == ".jsonl":
        return [json.loads(line) for line in decode_csv(path).splitlines() if line.strip()]
    if suffix == ".txt":
        return [{"review_text": line} for line in decode_csv(path).splitlines() if line.strip()]
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise ValueError("读取 XLSX 需要 openpyxl；请安装后重试，或另存为 CSV。") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(v or "") for v in next(values, [])]
        return [dict(zip(headers, row)) for row in values]
    raise ValueError(f"不支持的文件类型：{suffix or '无扩展名'}")


def build_mapping(headers: Iterable[str]) -> dict[str, str]:
    available = {norm_header(header): header for header in headers}
    mapping: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            if norm_header(alias) in available:
                mapping[field] = available[norm_header(alias)]
                break
    return mapping


def normalize_rating(value: Any) -> str:
    text = norm_text(value)
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        return ""
    rating = float(match.group())
    if rating < 0 or rating > 5:
        return text
    return str(int(rating)) if rating.is_integer() else str(rating)


def normalize_date(value: Any) -> str:
    text = norm_text(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def normalize_rows(rows: list[dict[str, Any]], default_platform: str) -> tuple[list[dict[str, str]], dict[str, Any]]:
    if not rows:
        return [], {"input_records": 0, "valid_records": 0, "unique_records": 0, "duplicates": 0}
    headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
    mapping = build_mapping(headers)
    if "review_text" not in mapping and not any("review_text" in row for row in rows):
        raise ValueError("未找到评论内容列。请使用 review/content/comment/评价内容/评论内容 等列名。")

    normalized: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    seen_fingerprints: set[str] = set()
    duplicate_reasons: Counter[str] = Counter()

    for row in rows:
        item = {field: norm_text(row.get(mapping.get(field, field), "")) for field in FIELDS if field != "duplicate_status"}
        item["platform"] = item["platform"] or default_platform
        item["rating"] = normalize_rating(item["rating"])
        item["date"] = normalize_date(item["date"])
        if not item["review_text"]:
            continue

        status = "unique"
        source_id = item["source_id"]
        fingerprint_raw = "|".join((dedupe_text(item["review_text"]), item["sku"].lower(), item["rating"]))
        fingerprint = hashlib.sha256(fingerprint_raw.encode("utf-8")).hexdigest()
        if source_id and source_id in seen_ids:
            status = "duplicate_source_id"
        elif fingerprint in seen_fingerprints:
            status = "duplicate_exact_text_sku_rating"
        if source_id:
            seen_ids.add(source_id)
        seen_fingerprints.add(fingerprint)
        item["duplicate_status"] = status
        if status != "unique":
            duplicate_reasons[status] += 1
        normalized.append({field: item.get(field, "") for field in FIELDS})

    unique_count = sum(1 for item in normalized if item["duplicate_status"] == "unique")
    summary = {
        "input_records": len(rows),
        "valid_records": len(normalized),
        "unique_records": unique_count,
        "duplicates": len(normalized) - unique_count,
        "duplicate_reasons": dict(duplicate_reasons),
        "mapped_fields": mapping,
        "platform_counts": dict(Counter(item["platform"] or "unknown" for item in normalized if item["duplicate_status"] == "unique")),
        "rating_counts": dict(Counter(item["rating"] or "unknown" for item in normalized if item["duplicate_status"] == "unique")),
        "sku_counts": dict(Counter(item["sku"] or "unknown" for item in normalized if item["duplicate_status"] == "unique")),
        "note": "脚本只做字段归一、精确去重和基础统计；主题、情绪及归因需基于证据另行分析。",
    }
    return normalized, summary


def write_outputs(output_dir: Path, rows: list[dict[str, str]], summary: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    with (output_dir / "normalized_reviews.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    (output_dir / "normalization_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="归一化电商评论导出文件并做精确去重。")
    parser.add_argument("input", type=Path, help="CSV/TSV/JSON/JSONL/TXT/XLSX 文件")
    parser.add_argument("--output-dir", type=Path, default=Path("normalized-output"))
    parser.add_argument("--platform", default="", help="源文件没有平台列时使用的默认平台")
    args = parser.parse_args()
    try:
        rows = read_rows(args.input)
        normalized, summary = normalize_rows(rows, args.platform)
        write_outputs(args.output_dir, normalized, summary)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
