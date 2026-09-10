#!/usr/bin/env python3

import argparse
import csv
import math
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate an Amazon replenishment and shipment plan from a CSV or XLSX file."
    )
    parser.add_argument("--input", required=True, help="Input CSV or XLSX file.")
    parser.add_argument("--output", required=True, help="Output CSV or XLSX file.")
    parser.add_argument(
        "--selected-only",
        action="store_true",
        help="Keep only rows where selected is truthy.",
    )
    return parser.parse_args()


def normalize_header(value: str) -> str:
    return value.strip()


def to_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def to_bool(value) -> bool:
    if value is None:
        return True
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if not text:
        return True
    return text in {"1", "true", "yes", "y", "selected"}


def ceil_multiple(value: int, multiple: int) -> int:
    if multiple <= 0:
        return value
    return int(math.ceil(value / multiple) * multiple)


def read_csv(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError("Input CSV has no header row.")
        headers = [normalize_header(name) for name in reader.fieldnames]
        rows = []
        for raw_row in reader:
            row = {}
            for header, value in zip(headers, raw_row.values()):
                row[header] = value
            rows.append(row)
    return headers, rows


def read_xlsx(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError("Input workbook is empty.")
    raw_headers = values[0]
    headers = [normalize_header(str(cell or "")) for cell in raw_headers]
    rows = []
    for raw_row in values[1:]:
        row = {}
        for index, header in enumerate(headers):
            row[header] = raw_row[index] if index < len(raw_row) else None
        rows.append(row)
    return headers, rows


def read_rows(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError("Input must be a CSV or XLSX file.")


def write_csv(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def write_xlsx(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)


def write_rows(path: Path, headers: List[str], rows: List[Dict[str, object]]) -> None:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        write_csv(path, headers, rows)
        return
    if suffix in {".xlsx", ".xlsm"}:
        write_xlsx(path, headers, rows)
        return
    raise ValueError("Output must be a CSV or XLSX file.")


def compute_row(row: Dict[str, object]) -> Dict[str, object]:
    current_stock = to_number(row.get("current_stock"))
    inbound_on_way = to_number(row.get("inbound_on_way"))
    daily_sales = to_number(row.get("daily_sales"))
    lead_time_days = to_number(row.get("lead_time_days"))
    safety_stock_days = to_number(row.get("safety_stock_days"))
    target_days_override = to_number(row.get("target_days_override"))
    carton_qty = int(to_number(row.get("carton_qty")))
    min_ship_qty = int(to_number(row.get("min_ship_qty")))

    target_coverage_days = (
        target_days_override
        if target_days_override > 0
        else lead_time_days + safety_stock_days
    )
    target_stock = daily_sales * target_coverage_days
    raw_need = target_stock - current_stock - inbound_on_way
    net_needed_units = max(0, math.ceil(raw_need))
    recommended_ship_units = net_needed_units

    if recommended_ship_units > 0 and carton_qty > 0:
        recommended_ship_units = ceil_multiple(recommended_ship_units, carton_qty)

    if recommended_ship_units > 0 and min_ship_qty > 0:
        recommended_ship_units = max(recommended_ship_units, min_ship_qty)
        if carton_qty > 0:
            recommended_ship_units = ceil_multiple(recommended_ship_units, carton_qty)

    notes = []
    if target_days_override > 0:
        notes.append("used target_days_override")
    if carton_qty > 0:
        notes.append(f"rounded to carton_qty={carton_qty}")
    if min_ship_qty > 0:
        notes.append(f"enforced min_ship_qty={min_ship_qty}")
    if recommended_ship_units == 0:
        notes.append("no shipment needed")

    computed = dict(row)
    computed["target_coverage_days"] = target_coverage_days
    computed["target_stock"] = round(target_stock, 2)
    computed["net_needed_units"] = net_needed_units
    computed["recommended_ship_units"] = recommended_ship_units
    computed["calculation_notes"] = "; ".join(notes)
    return computed


def main() -> None:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()

    headers, rows = read_rows(input_path)
    if "sku" not in headers:
        raise ValueError("Missing required column: sku")

    computed_rows = []
    for row in rows:
        is_selected = to_bool(row.get("selected"))
        if args.selected_only and not is_selected:
            continue
        computed_rows.append(compute_row(row))

    extra_headers = [
        "target_coverage_days",
        "target_stock",
        "net_needed_units",
        "recommended_ship_units",
        "calculation_notes",
    ]
    output_headers = list(headers)
    for header in extra_headers:
        if header not in output_headers:
            output_headers.append(header)

    write_rows(output_path, output_headers, computed_rows)
    print(f"Wrote {len(computed_rows)} rows to {output_path}")


if __name__ == "__main__":
    main()
